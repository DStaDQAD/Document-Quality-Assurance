"""The rows the PII parser derives from SEKI V.39: net positions and external debt; and the
Cadangan Devisa parser's name for V.9's total."""
from unittest import mock

import pytest

from publication_parsers import cadangan_devisa, pii
from table_model import TableData


def test_the_reserves_total_is_called_what_the_release_calls_it():
    # "posisi cadangan devisa … 146,5" found 'Cadangan Devisa Lainnya' (125,4) before 'Total'.
    sheet = TableData(title="V.9. Posisi Cadangan Devisa", unit="Juta USD",
                      row_labels=["Cadangan Devisa Lainnya", "Total"])
    sheet._data = {("Cadangan Devisa Lainnya", 2026, "Aug"): 125377.24,
                   ("Total", 2026, "Aug"): 146507.61}
    with mock.patch.object(cadangan_devisa, "parse_with_bi_reader", return_value=sheet):
        table = cadangan_devisa.parse(b"", "5.9")

    assert table.lookup_fuzzy("Cadangan Devisa", 2026, "Aug") == ("Cadangan Devisa", 146507.61)
    assert table.lookup_fuzzy("cadangan devisa lainnya", 2026, "Aug")[1] == 125377.24

_ROWS = {
    "Aset > Investasi Portofolio": 43485.39,
    "Kewajiban > Investasi Portofolio": 273103.82,
    "Kewajiban > Investasi Langsung > Instrumen Utang": 47148.36,
    "Kewajiban > Investasi Portofolio > Surat Utang": 209618.08,
    "Kewajiban > Investasi Lainnya": 196602.10,
    "Ke luar negeri > Instrumen utang": 51264.29,
}


def _table():
    table = TableData(title="V.39. Posisi Investasi Internasional Indonesia", unit="Juta USD",
                      row_labels=list(_ROWS))
    for label, value in _ROWS.items():
        table._data[(label, 2026, "Q2")] = value
    return table


def _ora_xlsx(month_cell):
    """The team's 'Tabel Cadev ORA': SDDS Official Reserve Assets, one sheet called 'sheet'."""
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ORA"      # openpyxl renames "Sheet" -> "sheet" to "sheet1" (case-insensitive clash)
    ws.title = "sheet"
    rows = [
        [None, None, None, None, None, "Lampiran 1"],
        ["SDDS Data Category and Component", None, "Unit Description", "Data Observation"],
        [None, None, None, "Date of Latest Observation", "Latest Data", "Latest Data -1"],
        ["Official Reserve Assets", None, "Millions of US$", month_cell, 146320.43, 146507.61],
        ["     ", "Foreign currency reserves", "Millions of US$", month_cell, 125187.37, 124668.19],
        [None, "Gold", "Millions of US$", month_cell, 11739.34, 12446.11],
    ]
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("month_cell", ["September/26", "Septembr/26", "Sep/2026"])
def test_the_teams_ora_table_reads_this_month_and_the_one_before(month_cell):
    # Oktober 2026: a check with ORA + Indimon came back 0/0/6 — no reader found a period header.
    table = cadangan_devisa.parse(_ora_xlsx(month_cell), "sheet")

    assert table.unit == "Juta USD"
    assert table.lookup_fuzzy("cadangan devisa", 2026, "Sep") == ("Cadangan Devisa", 146320.43)
    assert table.lookup_fuzzy("cadangan devisa", 2026, "Aug") == ("Cadangan Devisa", 146507.61)
    assert table.lookup_fuzzy("Gold", 2026, "Aug")[1] == 12446.11
    assert table.term_aliases["emas"] == "Cadangan Devisa > Gold"


def test_the_teams_indimon_table_reads_the_reserves_memorandum_in_usd():
    # The SEKI reader took Indimon's first data column for its labels ('930537.86…'), silently.
    import datetime
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tabel SP (indonesian version)"
    rows = [
        ["PERKEMBANGAN BESARAN MONETER"], ["(Dalam Miliar Rupiah)"],
        [None, "Items", 2026, None],
        ["No.", None, "Agst", "Sep*"],
        [None, None, datetime.datetime(2026, 8, 31), datetime.datetime(2026, 9, 30), datetime.datetime(2026, 7, 31)],
        [1, "Uang Primer", 1835112.12, 1892362.26, 1800000.0],
        [None, "Memorandum item:"],
        [None, "Cadangan Devisa (konsep IRFCL)    4)"],
        [None, "Dalam Juta USD", 146507.61, 146320.43, 145343.82],
    ]
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)

    table = cadangan_devisa.parse(buf.getvalue(), "Tabel SP (indonesian version)")

    assert table.unit == "Miliar Rp" and table.row_units == {"Cadangan Devisa": "Juta USD"}
    assert table.lookup_fuzzy("cadangan devisa", 2026, "Sep") == ("Cadangan Devisa", 146320.43)
    assert table.lookup_fuzzy("Uang Primer", 2026, "Aug")[1] == 1835112.12


def test_a_sheet_without_the_ora_layout_falls_back_to_the_generic_readers():
    import io
    import openpyxl
    from publication_parsers import PublicationParseError
    wb = openpyxl.Workbook()
    wb.active.title = "ORA"
    wb.active.title = "sheet"
    wb.active.append(["Bukan tabel ORA", 1, 2])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(PublicationParseError):
        cadangan_devisa.parse(buf.getvalue(), "sheet")


def test_a_component_with_both_sides_gets_a_hidden_net_row():
    table = _table()
    pii._add_net_rows(table)

    row = table.net_rows["investasi portofolio"]
    assert table._data[(row, 2026, "Q2")] == pytest.approx(43485.39 - 273103.82)
    assert row not in table.row_labels
    assert "investasi langsung" not in table.net_rows  # one side only


def test_external_debt_is_the_liabilities_debt_instruments_and_the_report_finds_it():
    # "posisi total utang luar negeri (ULN) Indonesia (posisi KFLN dalam instrumen utang) …
    # 453,4": its 'instrumen utang' reached the memorandum's 'Ke luar negeri > Instrumen utang'.
    table = _table()
    assert pii._add_sum_row(table, pii.ULN_ROW, pii.ULN_PARTS)
    table.term_aliases = dict(pii.TERM_ALIASES)

    label, value = table.lookup_fuzzy(
        "total utang luar negeri (ULN) Indonesia (posisi KFLN dalam instrumen utang)", 2026, "Q2")

    assert label == pii.ULN_ROW
    assert value == pytest.approx(453368.54)


def test_afln_and_kfln_are_the_totals_only_as_a_whole_claim():
    # PII Tw I-2026: "Posisi AFLN Indonesia menurun …" found no row; "AFLN sektor publik" (a
    # breakdown V.39 lacks) and a ratio's "total AFLN" must still find none.
    table = TableData(title="V.39", unit="Juta USD", row_labels=["Aset", "Kewajiban"])
    table._data = {("Aset", 2026, "Q1"): 556651.3, ("Kewajiban", 2026, "Q1"): 784283.89}
    table.label_aliases = {name: row for name, row in pii.LABEL_ALIASES.items()
                           if row in table.row_labels}

    assert table.lookup_fuzzy("Posisi AFLN Indonesia", 2026, "Q1")[0] == "Aset"
    assert table.lookup_fuzzy("Kewajiban Finansial Luar Negeri (KFLN)", 2026, "Q1")[0] == "Kewajiban"
    for other in ("AFLN sektor publik", "total AFLN", "posisi KFLN jangka panjang", "KFLN"):
        assert table.lookup_fuzzy(other, 2026, "Q1") == (None, None)


def test_a_component_named_by_its_side_abbreviation_is_that_side_row():
    # PII Tw IV-2025: "peningkatan posisi AFLN maupun KFLN investasi lainnya sebesar 233 juta dolar
    # AS dan 178 juta dolar AS" resolved to 'Aset > Investasi Lainnya', and the one-group guard
    # then read 'Aset' as a group the claim never named: Tidak Cukup Data on two report errors.
    from paired_verifier import _answers_one_group_of_it
    from structured_extractor import ExtractedFact, PeriodPoint

    rows = ["Aset", "Aset > Investasi Lainnya", "Kewajiban", "Kewajiban > Investasi Lainnya"]
    table = TableData(title="V.39", unit="Juta USD", row_labels=list(rows))
    table._data = {(row, 2025, "Q4"): 1000.0 + i for i, row in enumerate(rows)}
    table.term_aliases = dict(pii.TERM_ALIASES)
    table.label_aliases = {name: row for name, row in {**pii.LABEL_ALIASES, **pii.COMPONENT_ALIASES}.items()
                           if row in table.row_labels}

    for claim, row in (("AFLN investasi lainnya", "Aset > Investasi Lainnya"),
                       ("AIFLN investasi lainnya", "Aset > Investasi Lainnya"),
                       ("posisi KFLN investasi lainnya", "Kewajiban > Investasi Lainnya")):
        assert table.lookup_fuzzy(claim, 2025, "Q4")[0] == row
        fact = ExtractedFact(operation="diff", periods=[PeriodPoint(claim, 2025, "Q4")],
                             claimed_value=233.0, unit="juta dolar AS", context_quote=claim)
        assert not _answers_one_group_of_it(fact, [(row, 1.0)], table)
    # A sector V.39 does not have keeps finding no row.
    assert table.lookup_fuzzy("AFLN sektor lainnya", 2025, "Q4") == (None, None)


def test_npi_reads_still_in_surplus_as_a_level():
    # NPI Tw I-2026: "Investasi langsung tetap mencatatkan surplus" came back as a rise on
    # 3,15 -> 2,02 — still a surplus, a smaller one.
    import re
    from publication_parsers import npi

    sentence = "Investasi langsung tetap mencatatkan surplus sebagai cerminan dari persepsi positif investor"
    matched = [side for pattern, side in npi.LEVEL_PHRASES.items() if re.search(pattern, sentence, re.I)]
    assert matched == [("above", 0.0)]
    assert "transaksi finansial" in npi.LAYOUT_LEVELS


def test_ppkom_names_each_city_row_and_its_property_types():
    # PPKom's city tables put the city's number in column 1 and the city's own figures on its row;
    # the reader filed the property types under the number ('6 > Sewa').
    from publication_parsers import ppkom

    table = TableData(title="Tabel 2. Perkembangan Indeks Harga Properti Komersial", unit="",
                      row_labels=["6 > MEDAN", "6 > Sewa", "7 > SEMARANG", "7 > Sewa"])
    for label, value in zip(table.row_labels, (98.48, 108.63, 134.68, 118.02)):
        table._data[(label, 2026, "Q2")] = value
    ppkom._name_cities(table)

    assert table.row_labels == ["MEDAN", "MEDAN > Sewa", "SEMARANG", "SEMARANG > Sewa"]
    assert table._data[("SEMARANG > Sewa", 2026, "Q2")] == 118.02
