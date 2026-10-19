"""Which sheets each publication parser claims, and how the reader wrappers report failure."""
import re
from unittest.mock import patch

import pytest

from publication_parsers import PublicationParseError, covers, parse_for_publication
from table_model import TableData


@pytest.mark.parametrize("publication, sheet, expected", [
    ("uang-beredar", "I.1", True),
    ("uang-beredar", "I.1A", True),
    ("uang-beredar", " i.1a ", True),
    ("uang-beredar", "I.10", False),
    ("uang-beredar", "Th 1985-1992", False),
    ("uang-primer-m0", "I.2", True),
    ("uang-primer-m0", "Th 2010-2021", False),
    ("cadangan-devisa", "5.9", True),
    ("cadangan-devisa", "5.90", False),
    ("cadangan-devisa", "I.1", False),
    ("", "I.1", False),
    ("tidak-ada", "I.1", False),
])
def test_covers(publication, sheet, expected):
    assert covers(publication, sheet) is expected


def test_an_uncovered_sheet_is_left_to_the_generic_cascade():
    assert parse_for_publication("cadangan-devisa", b"never read", "I.1") is None


def test_seki_wrapper_reports_an_unreadable_file_as_a_publication_error():
    with pytest.raises(PublicationParseError, match="SEKI"):
        parse_for_publication("uang-beredar", b"not an excel file", "I.1")


@patch("publication_parsers._common.parse_bi_table")
def test_seki_wrapper_rejects_a_table_whose_hierarchy_collapsed(mock_bi):
    mock_bi.return_value = TableData(
        title="t", unit="Miliar Rp", row_labels=["Rupiah", "Rupiah"],
        _data={("Rupiah", 2026, "Jul"): 1.0},
    )

    with pytest.raises(PublicationParseError, match="berulang"):
        parse_for_publication("uang-beredar", b"bytes", "I.1")


@patch("publication_parsers._common.parse_bi_table")
def test_seki_wrapper_rejects_a_table_without_figures(mock_bi):
    mock_bi.return_value = TableData(title="t", unit="Miliar Rp", row_labels=["M2"])

    with pytest.raises(PublicationParseError, match="angka"):
        parse_for_publication("uang-primer-m0", b"bytes", "I.2")


@pytest.mark.parametrize("sheet, expected", [
    ("TabI.1", True), ("TabI.7", True), ("Tbl II.1", True), ("Tbl II.6", True),
    ("Tbl II.7", False), ("Tbl II.8", False),
    ("Tbl III.1", True), ("Tbl III.10", True), ("Tbl III.11", False),
])
def test_sulni_covers_only_its_time_series_sheets(sheet, expected):
    assert covers("sulni", sheet) is expected


@pytest.mark.parametrize("publication, sheet, expected", [
    ("npi", "5.1", True), ("npi", "Th 2004-2010", False), ("npi", "I.1", False),
    ("pii", "5.39", True), ("pii", "2001-2012", False), ("pii", "5.3", False),
])
def test_npi_and_pii_cover_their_seki_sheets(publication, sheet, expected):
    assert covers(publication, sheet) is expected


@pytest.mark.parametrize("publication, sheet, expected", [
    ("sk", "Tabel 1", True), ("sk", "Tabel 9", True), ("sk", "Tabel 10", False),
    ("skdu", "T1 Kegiatan Usaha", True), ("skdu", "T7b Investasi Semesteran", True),
    ("skdu", "T10 Margin Usaha", True), ("skdu", "Tabel 1", False),
])
def test_sk_and_skdu_cover_their_survey_sheets(publication, sheet, expected):
    assert covers(publication, sheet) is expected


@patch("publication_parsers._common.parse_generic_table")
def test_survey_wrapper_keeps_the_indonesian_half_of_the_unit(mock_generic):
    mock_generic.return_value = TableData(
        title="t", unit="Rp)  / (IDR", row_labels=["Mandor"], _data={("Mandor", 2026, "Q1"): 1.0},
    )

    assert parse_for_publication("skdu", b"bytes", "T9 Upah Rata-rata").unit == "Rp"


@patch("publication_parsers._common.parse_generic_table")
def test_survey_wrapper_fills_in_the_unit_a_consumer_survey_sheet_leaves_out(mock_generic):
    mock_generic.side_effect = lambda data, sheet: TableData(
        title="t", unit="", row_labels=["IKK"], _data={("IKK", 2026, "Aug"): 120.0},
    )

    assert parse_for_publication("sk", b"bytes", "Tabel 1").unit == "Indeks"
    assert parse_for_publication("sk", b"bytes", "Tabel 5").unit == "%"


@patch("publication_parsers._common.parse_generic_table")
def test_survey_wrapper_rejects_a_table_that_is_not_a_time_series(mock_generic):
    mock_generic.return_value = TableData(
        title="t", unit="", row_labels=["IKK"], axis_type="categorical",
        _data={("IKK", "Nilai"): 1.0},
    )

    with pytest.raises(PublicationParseError, match="deret waktu"):
        parse_for_publication("sk", b"bytes", "Tabel 1")


@pytest.mark.parametrize("sheet, expected", [
    ("Tabel 1", True), ("Tabel 4", True), ("Tabel 9", True), ("Tabel 10", False), ("Tabel1", False),
])
def test_spe_covers_its_nine_tables(sheet, expected):
    assert covers("spe", sheet) is expected


@pytest.mark.parametrize("sheet, expected", [
    ("Tabel1", True), ("Tabel 1", True), ("tabel4", True), ("Tabel 5 (disc)", True), ("Tabel6", False),
])
def test_sbank_covers_its_sheets_however_they_are_spaced(sheet, expected):
    assert covers("sbank", sheet) is expected


@pytest.mark.parametrize("sheet, expected", [
    ("TABEL 1", True), ("Tabel 2", True), ("TABEL 3", True), ("TABEL 4", False),
])
def test_shpr_covers_its_three_tables(sheet, expected):
    assert covers("shpr", sheet) is expected


@pytest.mark.parametrize("sheet, expected", [
    ("T1 PMI", True), ("T1 - Komponen PMI", True), ("T2", True), ("T2 Sublapangan", True),
    ("T10", False), ("T3", False),
])
def test_pmi_covers_both_sheet_namings(sheet, expected):
    assert covers("pmi", sheet) is expected


def test_every_publication_in_the_ui_has_its_own_parser():
    import check_history
    from publication_parsers import PARSERS

    assert sorted(PARSERS) == sorted(check_history.PUBLICATIONS)


def test_pmi_tables_say_their_index_is_read_in_percent(monkeypatch):
    # The report prints PMI-BI as "52,03%"; the verifier only compares a percent claim with an
    # index whose unit carries the % (see paired_verifier._is_plain_index).
    from publication_parsers import _common

    grid = [
        [None, "TABEL 1. PROMPT MANUFACTURING INDEX - BANK INDONESIA"],
        [None, "(%, Indeks)  / (%, Index)"],
        [None, "Komponen PMI", 2026],
        [None, "Komponen PMI", "I"],
        [None, "PMI - BI", 52.03125],
    ]
    monkeypatch.setattr(_common, "load_grid", lambda data, sheet: grid)

    assert parse_for_publication("pmi", b"bytes", "T1 PMI").unit == "%, Indeks"


@pytest.mark.parametrize("publication, sheet, unit", [
    ("spe", "Tabel 1", "Indeks"), ("spe", "Tabel 5", "Indeks"), ("spe", "Tabel 9", "Indeks"),
    ("spe", "Tabel 2", "%, yoy"), ("spe", "Tabel 4", "%, yoy"),
    ("spe", "Tabel 6", "%, yoy"), ("spe", "Tabel 8", "%, yoy"),
    ("spe", "Tabel 3", "%, mtm"), ("spe", "Tabel 7", "%, mtm"),
    ("shpr", "TABEL 1", "Indeks (2018=100)"), ("shpr", "TABEL 3", "%, qtq & yoy"),
])
def test_growth_tables_say_which_growth_they_hold(publication, sheet, unit):
    # The verifier reads a '%, yoy' cell as the answer to a yoy claim instead of computing growth
    # over it, and keeps a yoy claim off an mtm / qtq table (paired_verifier._growth_kind).
    from publication_parsers import PARSERS
    from publication_parsers._common import spec_for

    assert spec_for(sheet, PARSERS[publication].SPECS).unit == unit


def test_npi_adds_the_memo_nonmigas_balance_the_report_computes():
    # NPI Tw II-2026: "neraca perdagangan nonmigas surplus 11,5" = memo exports 70.471,04 plus
    # memo imports -58.970,41 (stored negative); the general-merchandise row says 12,1.
    from publication_parsers.npi import TERM_ALIASES, _add_memo_balances

    exports, imports = "Nominal > Total Ekspor (fob) > Nonmigas", "Nominal > Total Impor (fob) > Nonmigas"
    table = TableData(title="NPI", unit="Juta USD", row_labels=[exports, imports],
                      _data={(exports, 2026, "Q2"): 70471.04, (imports, 2026, "Q2"): -58970.41,
                             (exports, 2026, "Q1"): 63537.06})
    _add_memo_balances(table)

    name = TERM_ALIASES["neraca perdagangan nonmigas"]
    assert name in table.row_labels
    assert table.lookup(name, 2026, "Q2") == pytest.approx(11500.63)
    assert table.lookup(name, 2026, "Q1") is None  # no import figure to pair with


def test_npi_memo_balance_is_skipped_when_a_sheet_lacks_the_memo_rows():
    from publication_parsers.npi import _add_memo_balances

    table = TableData(title="5.1", unit="Juta USD", row_labels=["Barang > Ekspor"],
                      _data={("Barang > Ekspor", 2026, "Q2"): 1.0})
    _add_memo_balances(table)
    assert table.row_labels == ["Barang > Ekspor"]


def test_letter_spaced_total_is_read_as_total():
    from publication_parsers._common import _close_letter_spacing

    table = TableData(title="T7", unit="SBT", row_labels=["Konstruksi", "T O T A L"],
                      _data={("T O T A L", 2026, "Q2"): 5.46, ("Konstruksi", 2026, "Q2"): 0.26})
    _close_letter_spacing(table)
    assert table.row_labels == ["Konstruksi", "TOTAL"]
    assert table.lookup("TOTAL", 2026, "Q2") == 5.46


@patch("publication_parsers._common.parse_bi_table")
def test_uang_beredar_tables_name_the_report_sections_they_do_not_hold(mock_bi):
    mock_bi.return_value = TableData(
        title="t", unit="Miliar Rp", row_labels=["Simpanan Berjangka (Rupiah dan Valas)"],
        _data={("Simpanan Berjangka (Rupiah dan Valas)", 2026, "Aug"): 1.0},
    )

    table = parse_for_publication("uang-beredar", b"bytes", "I.1A")

    headings = {"PERKEMBANGAN DANA PIHAK KETIGA (DPK)": True, "PERKEMBANGAN KREDIT": True,
                "PERKEMBANGAN UANG PRIMER ADJUSTED": True, "SUKU BUNGA KREDIT DAN SIMPANAN": True,
                "KOMPONEN UANG BEREDAR": False, "FAKTOR-FAKTOR YANG MEMENGARUHI UANG BEREDAR": False}
    for heading, uncovered in headings.items():
        hit = any(re.search(p, heading, re.IGNORECASE) for p in table.uncovered_sections)
        assert hit == uncovered, heading


def test_npi_names_passenger_transport_with_its_parent():
    # NPI Tw II-2026: "defisit jasa transportasi penumpang" found no row, because the parser kept
    # the bare 'Penumpang' (unique) while 'Transportasi > Barang' got its parent (repeated name).
    from publication_parsers import npi

    table = TableData(title="t", unit="Juta USD",
                      row_labels=["Transportasi", "Penumpang", "Penumpang > Impor", "Transportasi > Barang"],
                      _data={("Penumpang", 2026, "Q2"): -80.98, ("Penumpang > Impor", 2026, "Q2"): -494.94,
                             ("Transportasi", 2026, "Q2"): -2502.89,
                             ("Transportasi > Barang", 2026, "Q2"): -2292.5})

    npi._add_parents(table)

    assert table.row_labels == ["Transportasi", "Transportasi > Penumpang",
                                "Transportasi > Penumpang > Impor", "Transportasi > Barang"]
    assert table._data[("Transportasi > Penumpang", 2026, "Q2")] == -80.98
    assert table.lookup_fuzzy("defisit jasa transportasi penumpang", 2026, "Q2") == (
        "Transportasi > Penumpang", -80.98)


def test_npi_reads_coupon_payments_as_portfolio_interest():
    from publication_parsers import npi

    assert npi.TERM_ALIASES["pembayaran kupon/bunga"] == (
        "Pendapatan investasi portofolio > Pendapatan utang (bunga) > Pembayaran")
    # Safe only with the income phrases beside it, now that aliases apply in one pass.
    assert npi.TERM_ALIASES["investasi lainnya"] == "Transaksi Finansial > Investasi Lainnya"
    assert npi.TERM_ALIASES["pendapatan atas investasi lainnya"] == "Pendapatan investasi lainnya"


def test_aliases_apply_in_one_pass_and_twice_change_nothing():
    # NPI: "pendapatan atas investasi lainnya" must stay the income row although the shorter
    # "investasi lainnya" points to the financial-account row.
    table = TableData(title="t", unit="Juta USD", row_labels=[])
    table.term_aliases = {"pendapatan atas investasi lainnya": "Pendapatan investasi lainnya",
                          "pendapatan investasi lainnya": "Pendapatan investasi lainnya",
                          "investasi lainnya": "Transaksi Finansial > Investasi Lainnya"}

    once = table._apply_aliases("neraca pendapatan atas investasi lainnya")

    assert once == "neraca Pendapatan investasi lainnya"
    assert table._apply_aliases(once) == once
    assert table._apply_aliases("Investasi Lainnya") == "Transaksi Finansial > Investasi Lainnya"


@patch("publication_parsers._common.parse_bi_table")
def test_uang_beredar_lists_the_m0_items_it_does_not_hold(mock_bi):
    mock_bi.return_value = TableData(title="t", unit="Miliar Rp", row_labels=["Uang Kuasi"],
                                     _data={("Uang Kuasi", 2026, "Aug"): 1.0})

    table = parse_for_publication("uang-beredar", b"bytes", "I.1A")

    for name in ("uang kartal yang diedarkan", "giro bank umum di Bank Indonesia adjusted", "Uang Primer (M0) adjusted"):
        assert any(re.search(p, name, re.IGNORECASE) for p in table.absent_items), name
    assert not any(re.search(p, "Uang Kartal di Luar Bank Umum dan BPR", re.IGNORECASE) for p in table.absent_items)


def test_pmi_reads_the_internal_t2_layout_without_the_short_name_column():
    # The SharePoint release workbook (Tw II/III-2026) has English in column 3 and Indonesian in
    # column 4, no column 2: the public layout's spec found no labels at all.
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "T2 PMI - Sublapangan Usaha"
    ws.cell(2, 4, "TABEL 2. PROMPT MANUFACTURING INDEX - BANK INDONESIA (PMI-BI)")
    ws.cell(3, 4, "(%, Indeks)  / (Index, %)")
    ws.cell(5, 4, "Subsector")
    ws.cell(5, 5, "Sublapangan Usaha Industri Pengolahan")
    ws.cell(5, 6, 2026)
    for col, quarter in ((6, "I"), (7, "II"), (8, "III*")):
        ws.cell(6, col, quarter)
    for row, (english, indonesian, values) in enumerate((
            ("Manufacture of Food Products and Beverages", "Industri Makanan dan Minuman", (51.2, 52.4, 53.1)),
            ("Manufacture of Tobacco Products", "Industri Pengolahan Tembakau", (49.8, 50.6, 51.0))), start=7):
        ws.cell(row, 4, english)
        ws.cell(row, 5, indonesian)
        for col, value in zip((6, 7, 8), values):
            ws.cell(row, col, value)
    buffer = io.BytesIO()
    wb.save(buffer)

    table = parse_for_publication("pmi", buffer.getvalue(), "T2 PMI - Sublapangan Usaha")

    assert table.row_labels == ["Industri Makanan dan Minuman", "Industri Pengolahan Tembakau"]
    assert table._data[("Industri Makanan dan Minuman", 2026, "Q2")] == 52.4


def test_pmi_reads_sector_performance_as_the_overall_index():
    from publication_parsers import pmi

    assert pmi.TERM_ALIASES["kinerja lu industri pengolahan"] == "PMI - BI"
    # SKDU's wording in the same release is a different survey.
    assert not any("kegiatan" in phrase for phrase in pmi.TERM_ALIASES)


def test_sbank_marks_its_net_balance_tables_and_names_dpk(monkeypatch):
    # "Penyaluran kredit baru diprakirakan tetap tumbuh dengan SBT 84,65%" is a positive balance;
    # the verifier reads it so only when the unit says SBT (paired_verifier._NET_BALANCE_UNIT).
    from publication_parsers import PARSERS, _common
    from publication_parsers._common import spec_for

    specs = PARSERS["sbank"].SPECS
    assert spec_for("Tabel1", specs).unit == "SBT, %"
    assert spec_for("Tabel3", specs).unit == "SBT, %"
    assert spec_for("Tabel4", specs).unit == "%"  # interest rates

    grid = [
        [None, "Tabel 3. Prakiraan Arah Pertumbuhan DPK (SBT, %)"],
        [None, "Periode", "Kelompok Bank", "Jenis Simpanan", "Period", 2026],
        [None, None, None, None, None, "III"],
        [None, "Prakiraan per Triwulan", "Gabungan", "Deposito", "Estimation", 40.97],
        [None, None, None, "Total", None, 87.52],
    ]
    monkeypatch.setattr(_common, "load_grid", lambda data, sheet: grid)

    table = parse_for_publication("sbank", b"bytes", "Tabel3")
    assert table.term_aliases["dpk"] == "Gabungan > Total"
    assert "gabungan" in table.layout_levels
    # Not on Tabel1, whose 'TOTAL' is new loans.
    monkeypatch.setattr(_common, "load_grid", lambda data, sheet: [
        [None, "Tabel 1. Penyaluran Kredit Baru (SBT, %)"],
        [None, "Jenis Kredit", "Rincian Kredit", "Type", 2026],
        [None, None, None, None, "II"],
        [None, "TOTAL", None, "TOTAL", 93.08],
    ])
    assert parse_for_publication("sbank", b"bytes", "Tabel1").term_aliases == {}


def test_only_sbank_skips_its_methodology_section():
    from publication_parsers import skipped_sections

    assert skipped_sections("sbank") == (r"metodologi",)
    assert skipped_sections("npi") == ()
    assert skipped_sections("") == ()


def test_skdu_marks_its_rows_tables_and_measure(monkeypatch):
    from publication_parsers import skdu

    def reader(rows, unit):
        return lambda data, sheet: TableData(title="t", unit=unit, row_labels=list(rows))

    monkeypatch.setattr(skdu, "parse_with_generic_reader", reader(
        ["Mudah", "Sulit", "Saldo Bersih (% Mudah - % Sulit)"], "% Responden"))
    access = skdu.parse(b"", "T3 Keuangan dan Akses Kredit")
    assert access.term_aliases["akses kredit"] == "Saldo Bersih (% Mudah - % Sulit)"
    assert not access.measure_levels  # its rows ARE balances by name

    monkeypatch.setattr(skdu, "parse_with_generic_reader", reader(
        ["PERTANIAN, KEHUTANAN DAN PERIKANAN", "TOTAL"], "% Saldo Bersih Tertimbang - SBT"))
    activity = skdu.parse(b"", "T1 Kegiatan Usaha")
    assert "kinerja" in activity.subject_words
    assert "saldo bersih" in activity.measure_levels
    assert not skdu.parse(b"", "T4 Penggunaan Tenaga Kerja").subject_words


def test_skdu_tabel_3_section_headings_are_layout(monkeypatch):
    from publication_parsers import skdu

    monkeypatch.setattr(skdu, "parse_with_generic_reader", lambda data, sheet: TableData(
        title="t", unit="% Responden", row_labels=["Kondisi keuangan selama 3 bulan terakhir : > - Likuiditas > Baik"]))

    assert "kondisi keuangan selama 3 bulan terakhir :" in skdu.parse(b"", "T3 Keuangan dan Akses Kredit").layout_levels
    assert not skdu.parse(b"", "T1 Kegiatan Usaha").layout_levels


def test_skdu_tabel_6_holds_yoy_rates(monkeypatch):
    # "prakiraan inflasi … 2,61% (yoy)" is the cell itself, not a growth to compute over it.
    from publication_parsers import skdu

    monkeypatch.setattr(skdu, "parse_with_generic_reader",
                        lambda data, sheet: TableData(title="t", unit="%", row_labels=["TOTAL"]))

    assert skdu.parse(b"", "T6 Prakiraan Inflasi").unit == "%, yoy"
    assert skdu.parse(b"", "T2 Kapasitas Produksi").unit == "%"
